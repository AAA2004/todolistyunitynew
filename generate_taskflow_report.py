from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement


WORKSPACE = Path("/Users/ahmedashraf/todolistapp")
TEMPLATE = Path("/Users/ahmedashraf/Downloads/Tutorial 3 - Report Template (1).docx")
USE_CASE_IMG = Path("/Users/ahmedashraf/Downloads/use case todolist.png")
ACTIVITY_IMG = Path("/Users/ahmedashraf/Downloads/activity diagram to do list.png")
OUTPUT = WORKSPACE / "TaskFlow_Report.docx"


def remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)
    paragraph._p = paragraph._element = None


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if style:
        new_paragraph.style = style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def clear_section(doc: Document, heading_text: str, next_heading_text: str | None) -> None:
    paragraphs = doc.paragraphs
    start_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == heading_text)
    if next_heading_text is None:
        end_idx = len(paragraphs)
    else:
        end_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == next_heading_text)
    for idx in range(end_idx - 1, start_idx, -1):
        remove_paragraph(doc.paragraphs[idx])


def add_body_paragraph(anchor: Paragraph, text: str) -> Paragraph:
    return insert_paragraph_after(anchor, text, "Body Text")


def add_numbered_paragraph(anchor: Paragraph, text: str) -> Paragraph:
    p = insert_paragraph_after(anchor, text)
    try:
        p.style = "List Number"
    except KeyError:
        p.style = "List Paragraph"
    return p


def add_bulleted_paragraph(anchor: Paragraph, text: str) -> Paragraph:
    p = insert_paragraph_after(anchor, text)
    try:
        p.style = "List Bullet"
    except KeyError:
        p.style = "List Paragraph"
    return p


def add_heading3(anchor: Paragraph, text: str) -> Paragraph:
    return insert_paragraph_after(anchor, text, "Heading 3")


def add_caption(anchor: Paragraph, text: str) -> Paragraph:
    p = insert_paragraph_after(anchor, text, "Body Text")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.italic = True
    return p


def add_picture(anchor: Paragraph, image_path: Path, width_inches: float) -> Paragraph:
    p = insert_paragraph_after(anchor)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    return p


def find_heading(doc: Document, text: str) -> Paragraph:
    return next(p for p in doc.paragraphs if p.text.strip() == text)


def count_words(paragraphs: list[str]) -> int:
    words = []
    for text in paragraphs:
        words.extend(re.findall(r"\b[\w'-]+\b", text))
    return len(words)


def main() -> None:
    doc = Document(str(TEMPLATE))

    section_order = [
        "Introduction",
        "Project Background",
        "Motivation",
        "Project Goals",
        "Another Application",
        "Proposed Application Modeling",
        "Product Database",
        "Resources",
    ]

    for idx, heading in enumerate(reversed(section_order)):
        next_heading = None if idx == 0 else list(reversed(section_order))[idx - 1]
        clear_section(doc, heading, next_heading)

    word_bank: list[str] = []

    intro = find_heading(doc, "Introduction")
    intro_texts = [
        (
            "TaskFlow is a mobile to-do list and reminder application developed with Unity and C#. "
            "The main purpose of the app is to help users organise daily tasks in a simple and visual way. "
            "Instead of writing tasks on paper or in many separate notes, the user can keep all important "
            "activities in one place, grouped by day of the week. The app also reminds the user about "
            "upcoming tasks, which supports better time management and reduces the chance of forgetting "
            "important work."
        ),
        (
            "This project is important because many students and workers need a lightweight planning tool "
            "that is easy to understand and fast to use. TaskFlow gives the user the ability to sign up, "
            "log in, add or edit tasks, set priorities, assign tasks to one or more days, mark tasks as "
            "completed, and check weekly progress. The final output of the project is a functional cross-platform "
            "mobile app with local data storage and Firebase synchronisation."
        ),
    ]
    anchor = intro
    for text in intro_texts:
        anchor = add_body_paragraph(anchor, text)
        word_bank.append(text)

    background = find_heading(doc, "Project Background")
    background_texts = [
        (
            "TaskFlow is designed as a mobile application for Android and iOS. The project uses Unity as the "
            "development environment, C# for scripting, UI Toolkit for the interface, and Firebase Realtime "
            "Database for cloud storage. On the device side, the app can also keep a local copy of data so the "
            "user can still open the app and review tasks when the internet connection is weak or not available."
        ),
        (
            "From a hardware point of view, the app only needs a modern smartphone or tablet with basic internet "
            "access. From a software point of view, the project requires Unity, Firebase configuration files, and "
            "the mobile build tools for the chosen platform. The development process starts with interface design, "
            "then task data modelling, then user account handling, and finally reminder and progress features."
        ),
    ]
    anchor = background
    for text in background_texts:
        anchor = add_body_paragraph(anchor, text)
        word_bank.append(text)

    motivation = find_heading(doc, "Motivation")
    motivation_intro = (
        "The motivation behind TaskFlow comes from a common problem: many users know what they need to do, "
        "but they struggle to organise tasks clearly and follow them during a busy week. A to-do app becomes "
        "more useful when it does not only store tasks, but also gives structure, reminders, and visible progress. "
        "During development, the following challenges are expected:"
    )
    anchor = add_body_paragraph(motivation, motivation_intro)
    word_bank.append(motivation_intro)
    motivation_items = [
        "Keeping the interface simple while still supporting many features. This can be solved by dividing the app into clear pages such as login, day view, task details, and progress view.",
        "Managing user data correctly across local storage and cloud storage. This can be reduced by using clear data keys, regular sync rules, and testing save and load functions early.",
        "Creating reminders that are useful but not annoying. This can be improved by checking upcoming tasks at fixed times and showing only relevant popup messages.",
        "Supporting stable login and sign up behaviour. This challenge can be handled by validating input, checking stored credentials carefully, and using Firebase as a trusted backend service.",
    ]
    for text in motivation_items:
        anchor = add_numbered_paragraph(anchor, text)
        word_bank.append(text)
    motivation_close = (
        "These challenges can be reduced by following good design practice and testing each feature step by step."
    )
    anchor = add_body_paragraph(anchor, motivation_close)
    word_bank.append(motivation_close)

    goals = find_heading(doc, "Project Goals")
    goals_intro = (
        "The main goal of TaskFlow is to give the user a reliable and easy way to plan the week. "
        "The user should be able to achieve the following goals in the application:"
    )
    anchor = add_body_paragraph(goals, goals_intro)
    word_bank.append(goals_intro)
    goal_items = [
        "Create a new account and log in securely.",
        "Log out of the app when needed.",
        "Choose any day of the week and view the tasks linked to that day.",
        "Add a new task with a name, description, time, and priority.",
        "Assign one task to one day or to several days.",
        "Mark finished tasks as completed and see weekly progress.",
        "Receive reminder popups before important tasks are due.",
    ]
    for text in goal_items:
        anchor = add_bulleted_paragraph(anchor, text)
        word_bank.append(text)
    goals_close = (
        "A secondary goal is to keep the app fast and practical for everyday use. The user should not need many "
        "screens or complex steps to perform basic actions."
    )
    anchor = add_body_paragraph(anchor, goals_close)
    word_bank.append(goals_close)

    another = find_heading(doc, "Another Application")
    another_texts = [
        (
            "One well-known application with a similar purpose is Todoist. Todoist also helps users create tasks, "
            "organise schedules, and track progress. It is successful because it offers a clean design and strong task "
            "management features across many devices."
        ),
        (
            "TaskFlow is different because it is built around a simple weekly structure and a focused reminder flow. "
            "Instead of giving many advanced options at the start, it guides the user through a direct process: choose "
            "a day, open the tasks for that day, and manage them quickly. The app also combines local saving with Firebase "
            "sync, which supports both convenience and backup."
        ),
    ]
    anchor = another
    for text in another_texts:
        anchor = add_body_paragraph(anchor, text)
        word_bank.append(text)

    modeling = find_heading(doc, "Proposed Application Modeling")
    modeling_intro = (
        "The proposed system model shows how the user interacts with TaskFlow and how the main features support the full planning journey. "
        "The use case diagram presents the actors and features, while the activity diagram explains the order of actions from app start until logout."
    )
    anchor = add_body_paragraph(modeling, modeling_intro)
    word_bank.append(modeling_intro)

    anchor = add_heading3(anchor, "Use Case Diagram")
    use_case_text = (
        "Figure 1 shows the complete use case diagram for the application. The main actor is the user, while Firebase Database is the supporting system actor. "
        "The diagram includes sign up, login, logout, selecting the day of the week, adding or editing a task, assigning a task to days, saving data locally, "
        "syncing data with Firebase, marking a task as completed, viewing weekly progress, and receiving task reminders."
    )
    anchor = add_body_paragraph(anchor, use_case_text)
    word_bank.append(use_case_text)
    anchor = add_picture(anchor, USE_CASE_IMG, 6.5)
    anchor = add_caption(anchor, "Figure 1. Full use case diagram for TaskFlow.")

    anchor = add_heading3(anchor, "Use Case Scenario 1: User Login")
    scenario1 = [
        "Use case name: User Login.",
        "Primary actor: User.",
        "Goal: The user wants to access personal tasks and progress data.",
        "Preconditions: The user has already created an account, and the app is open on the login page.",
        "Trigger: The user enters username and password, then presses the login button.",
        "Main success scenario: First, the app checks whether the entered information matches the saved credentials. Next, the app verifies the data with the backend if needed. After successful validation, the app loads the user task data and opens the weekly day selection page.",
        "Alternative or exception flows: If the username or password is empty, the app shows an error message. If the credentials are incorrect, the app denies access and asks the user to try again. If the internet is not available, the app can still try local data.",
        "Postconditions: The user is logged in and can view or manage tasks.",
    ]
    for text in scenario1:
        anchor = add_body_paragraph(anchor, text)
        word_bank.append(text)

    anchor = add_heading3(anchor, "Use Case Scenario 2: Add or Edit Task")
    scenario2 = [
        "Use case name: Add or Edit Task.",
        "Primary actor: User.",
        "Goal: The user wants to create a new task or update an existing one for better planning.",
        "Preconditions: The user is logged in and has opened a selected day page.",
        "Trigger: The user chooses the add button or selects an existing task from the daily list.",
        "Main success scenario: The app opens the task details page. The user enters or updates the task name, description, time, and priority, then chooses one or more days. After pressing save, the app stores the data locally and syncs it with Firebase when the internet is available. The updated task then appears in the correct day list.",
        "Alternative or exception flows: If the task name is empty, the app does not save the task. If there is no connection, the app still keeps the task locally for later synchronisation. If the user edits a task, the app replaces the old values.",
        "Postconditions: A new task is created or an existing task is updated, and the schedule reflects the latest data.",
    ]
    for text in scenario2:
        anchor = add_body_paragraph(anchor, text)
        word_bank.append(text)

    anchor = add_heading3(anchor, "Activity Diagram")
    activity_text = (
        "Figure 2 presents the activity diagram for the app flow. The process starts when the application opens. If the user is not logged in, the system moves to sign up or login and checks the credentials. "
        "After that, the user selects a day and views the tasks for that day. From this point, the user can add or edit tasks, mark tasks as completed, view weekly progress, or wait for reminder checks. "
        "The activity diagram also shows the save process: data is stored locally first, then synced with Firebase if the internet is available."
    )
    anchor = add_body_paragraph(anchor, activity_text)
    word_bank.append(activity_text)
    anchor = add_picture(anchor, ACTIVITY_IMG, 5.8)
    anchor = add_caption(anchor, "Figure 2. Activity diagram for the TaskFlow application flow.")
    modeling_close = (
        "Together, the use case and activity diagrams give a clear picture of how the system works and help the developer check the final implementation."
    )
    anchor = add_body_paragraph(anchor, modeling_close)
    word_bank.append(modeling_close)

    database = find_heading(doc, "Product Database")
    database_texts = [
        (
            "TaskFlow needs a backend because user information and task data should be stored safely and retrieved again after login. "
            "In this project, Firebase Realtime Database is a suitable backend because it is easy to integrate with Unity and it supports simple cloud synchronisation. "
            "The backend stores the user account record and task-related fields, while the mobile device keeps a local copy."
        ),
        (
            "The backend is needed for three main reasons. First, it protects user continuity because tasks are not lost when the user changes device or reinstalls the app. "
            "Second, it supports login and sign up by keeping account credentials in a central place. Third, it allows the app to restore data when the user comes back after some time, "
            "which improves reliability."
        ),
        (
            "The stored data includes user credentials such as username and password, task lists for each day, task descriptions, priority levels, task times, task state values such as pending or completed, "
            "and notification status values used to avoid repeated reminder popups. User progress is also represented through the completed task states."
        ),
        (
            "Although the template mentions product information, in this project the equivalent data is task information. Each task acts like a managed item inside the system. "
            "When the user saves a task, the app first writes the data locally, then checks internet availability, and finally syncs the same data to Firebase. When the user logs in, the app can retrieve the stored cloud data "
            "and rebuild the local view. This design supports usability and backup."
        ),
    ]
    anchor = database
    for text in database_texts:
        anchor = add_body_paragraph(anchor, text)
        word_bank.append(text)

    resources = find_heading(doc, "Resources")
    resources_intro = (
        "The following resources are the most useful for this project because they are either official technical references or strong guides for the same type of application design."
    )
    anchor = add_body_paragraph(resources, resources_intro)
    word_bank.append(resources_intro)
    resource_items = [
        "Unity Technologies. (2024). UI Toolkit manual. Unity Documentation.",
        "Firebase. (2024). Firebase Realtime Database for Unity. Firebase Documentation.",
        "Firebase. (2024). Add Firebase to your Unity project. Firebase Documentation.",
        "Microsoft. (2024). C# documentation for collections, DateTime, and asynchronous programming. Microsoft Learn.",
        "Visual Paradigm. (2024). UML use case diagram and activity diagram guides. Visual Paradigm Online.",
    ]
    for text in resource_items:
        anchor = add_body_paragraph(anchor, text)
        word_bank.append(text)

    total_words = count_words(word_bank)
    print(f"Estimated report word count: {total_words}")
    doc.save(str(OUTPUT))
    print(f"Saved report to: {OUTPUT}")


if __name__ == "__main__":
    main()
